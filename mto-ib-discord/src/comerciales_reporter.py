"""
Reporte mensual automático a comerciales y colaboradores.

Se ejecuta el día 1 de cada mes a las 09:15 Madrid.

Antes de ejecutar descarga el Excel 'Control comerciales.xlsx' de Dropbox
y lee la pestaña 'Comerciales' para obtener la lista actualizada de personas.
Si falla la descarga, usa el fallback de config.yaml.

- Comercial (Colaborador=NO): email con resumen de canjes del mes anterior
- Comercial (Colaborador=SI): email de liquidación completo
"""

from __future__ import annotations

import asyncio
import calendar
import io
import json
import smtplib
import ssl
from datetime import datetime, timezone
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from .email_utils import build_message as _build_email, FIRMA_HTML
from typing import List, Optional
from zoneinfo import ZoneInfo

import aiohttp
import stripe as _stripe
from loguru import logger

MADRID_ZONE = ZoneInfo("Europe/Madrid")

MESES_ES = [
    "", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]

# Dropbox Business: namespace del equipo "Easy Tax Advice" (root_namespace_id)
_DROPBOX_PATH_ROOT = json.dumps({".tag": "namespace_id", "namespace_id": "2239051875"})

# Ruta del Excel en Dropbox
_EXCEL_DROPBOX_PATH = (
    "/Easy Tax Advice/CLIENTES/74141 MTO OPCIONES, S.L"
    "/COMPARTIDA MTO/ADMINISTRACION/Comerciales/Control comerciales.xlsx"
)
_EXCEL_SHEET = "Comerciales"


class ComercialReporter:

    def __init__(self, cfg: dict):
        stripe_cfg  = cfg.get("stripe",    {})
        smtp_cfg    = cfg.get("smtp",      {})
        discord_cfg = cfg.get("discord",   {})
        logbook_cfg = cfg.get("logbook",   {})
        com_cfg     = cfg.get("comerciales", {})

        _stripe.api_key = stripe_cfg.get("api_key", "")

        self.smtp_host = smtp_cfg.get("host",     "mail.mtoopciones.com")
        self.smtp_port = int(smtp_cfg.get("port", 587))
        self.smtp_user = smtp_cfg.get("user",     "info@mtoopciones.com")
        self.smtp_pass = smtp_cfg.get("password", "")

        self.discord_token    = discord_cfg.get("bot_token",  "")
        self.discord_guild_id = discord_cfg.get("guild_id",   "1387903106439839886")
        self.log_webhook      = discord_cfg.get("log_webhook", "")

        # Credenciales Dropbox (reutilizadas del logbook)
        self.dropbox_app_key      = logbook_cfg.get("dropbox_app_key",      "")
        self.dropbox_app_secret   = logbook_cfg.get("dropbox_app_secret",   "")
        self.dropbox_refresh_token = logbook_cfg.get("dropbox_refresh_token", "")

        # Fallback si falla la descarga del Excel
        self.personas_fallback = com_cfg.get("personas", [])
        self.gratuitos         = int(com_cfg.get("gratuitos", 8))

    # ── Arranque ───────────────────────────────────────────────

    async def start(self) -> None:
        asyncio.ensure_future(self._monthly_loop())
        now   = datetime.now(MADRID_ZONE)
        fire  = self._next_fire(now)
        delta = (fire - now).total_seconds() / 3600
        logger.info(
            f"Comerciales reporter iniciado — próximo envío el "
            f"{fire.day:02d}/{fire.month:02d}/{fire.year} 09:15 Madrid ({delta:.1f}h)"
        )

    def _next_fire(self, now: datetime) -> datetime:
        if now.day == 1 and (now.hour < 9 or (now.hour == 9 and now.minute < 15)):
            return now.replace(hour=9, minute=15, second=0, microsecond=0)
        if now.month == 12:
            return now.replace(year=now.year + 1, month=1, day=1,
                               hour=9, minute=15, second=0, microsecond=0)
        return now.replace(month=now.month + 1, day=1,
                           hour=9, minute=15, second=0, microsecond=0)

    async def _monthly_loop(self) -> None:
        while True:
            try:
                now  = datetime.now(MADRID_ZONE)
                fire = self._next_fire(now)
                await asyncio.sleep((fire - now).total_seconds())
                await self.run_monthly_report()
            except asyncio.CancelledError:
                break
            except Exception as e:
                logger.error(f"ComercialReporter: error en monthly loop: {e}")
                await asyncio.sleep(3600)

    # ── Pipeline principal ──────────────────────────────────────

    async def run_monthly_report(self) -> None:
        now = datetime.now(MADRID_ZONE)
        year  = now.year  if now.month > 1 else now.year - 1
        month = now.month - 1 if now.month > 1 else 12
        mes_str = MESES_ES[month]

        logger.info(f"ComercialReporter: generando reportes para {mes_str} {year}…")

        # 1. Cargar lista de personas desde el Excel (con fallback a config)
        personas = await self._load_personas_from_excel()
        logger.info(f"ComercialReporter: {len(personas)} personas cargadas")

        # 2. Conteo de miembros Discord
        member_count = await self._get_discord_member_count()
        logger.info(f"ComercialReporter: miembros Discord = {member_count}")

        loop    = asyncio.get_event_loop()
        results = []

        for persona in personas:
            nombre      = persona.get("nombre", "")
            mail        = persona.get("mail",   "")
            codigo      = persona.get("codigo", "")
            colaborador = str(persona.get("colaborador", "false")).lower() in ("true", "si", "yes", "1")

            if not mail or not codigo:
                results.append({"nombre": nombre, "mail": mail, "tipo": "—",
                                 "uses": 0, "ok": False, "error": "Sin mail o código"})
                continue

            uses = await loop.run_in_executor(
                None, lambda c=codigo: self._get_coupon_uses_in_month(c, year, month)
            )
            logger.info(f"ComercialReporter: {nombre} ({codigo}) → {uses} usos en {mes_str}")

            tipo   = "Liquidación" if colaborador else "Resumen canjes"
            ok, err = True, ""
            try:
                if colaborador:
                    await self._send_colaborador_email(persona, uses, member_count, year, month, now)
                else:
                    await self._send_comercial_email(persona, uses, year, month)
            except Exception as e:
                ok, err = False, str(e)
                logger.error(f"ComercialReporter: error enviando a {nombre}: {e}")

            results.append({"nombre": nombre, "mail": mail, "tipo": tipo,
                             "uses": uses, "ok": ok, "error": err})

        logger.info("ComercialReporter: todos los reportes procesados ✅")
        await self._log_to_discord(results, mes_str, year, member_count)

    # ── Lectura del Excel desde Dropbox ─────────────────────────

    async def _load_personas_from_excel(self) -> List[dict]:
        """
        Descarga el Excel desde Dropbox, parsea la pestaña Comerciales
        y devuelve lista de dicts {nombre, mail, codigo, colaborador}.
        Si falla, usa el fallback de config.yaml.
        """
        try:
            token = await self._get_dropbox_token()
            if not token:
                raise RuntimeError("No se pudo obtener token Dropbox")

            async with aiohttp.ClientSession() as s:
                async with s.post(
                    "https://content.dropboxapi.com/2/files/download",
                    headers={
                        "Authorization":        f"Bearer {token}",
                        "Dropbox-API-Arg":      f'{{"path": "{_EXCEL_DROPBOX_PATH}"}}',
                        "Dropbox-API-Path-Root": _DROPBOX_PATH_ROOT,
                    },
                    timeout=aiohttp.ClientTimeout(total=30),
                ) as resp:
                    if resp.status != 200:
                        text = await resp.text()
                        raise RuntimeError(f"Dropbox HTTP {resp.status}: {text[:100]}")
                    data = await resp.read()

            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            if _EXCEL_SHEET not in wb.sheetnames:
                raise RuntimeError(f"Pestaña '{_EXCEL_SHEET}' no encontrada en el Excel")

            ws = wb[_EXCEL_SHEET]
            personas = []
            header   = None

            for row in ws.iter_rows(values_only=True):
                values = [str(c).strip() if c is not None else "" for c in row]
                if not any(values):
                    continue
                if header is None:
                    # Primera fila con contenido = cabecera
                    header = [v.lower() for v in values]
                    continue

                def col(name: str) -> str:
                    try:
                        return values[header.index(name)]
                    except (ValueError, IndexError):
                        return ""

                nombre      = col("nombre")
                mail        = col("mail")
                codigo      = col("codigo")
                comercial   = col("comercial").upper()
                colaborador = col("colaborador").upper()

                if comercial != "SI" or not nombre:
                    continue

                personas.append({
                    "nombre":      nombre,
                    "mail":        mail,
                    "codigo":      codigo,
                    "colaborador": colaborador == "SI",
                })

            wb.close()
            logger.info(
                f"ComercialReporter: Excel descargado de Dropbox — "
                f"{len(personas)} comerciales encontrados"
            )
            return personas

        except Exception as e:
            logger.warning(
                f"ComercialReporter: no se pudo leer Excel de Dropbox ({e}), "
                f"usando fallback de config.yaml"
            )
            return self.personas_fallback

    async def _get_dropbox_token(self) -> Optional[str]:
        """Obtiene un access token de Dropbox usando el refresh token."""
        if not self.dropbox_refresh_token:
            return None
        try:
            async with aiohttp.ClientSession() as s:
                async with s.post(
                    "https://api.dropbox.com/oauth2/token",
                    data={
                        "grant_type":    "refresh_token",
                        "refresh_token": self.dropbox_refresh_token,
                        "client_id":     self.dropbox_app_key,
                        "client_secret": self.dropbox_app_secret,
                    },
                    timeout=aiohttp.ClientTimeout(total=15),
                ) as resp:
                    data = await resp.json()
                    return data.get("access_token")
        except Exception as e:
            logger.error(f"ComercialReporter: error obteniendo token Dropbox: {e}")
            return None

    # ── Stripe: canjes del mes ──────────────────────────────────

    def _get_coupon_uses_in_month(self, coupon_id: str, year: int, month: int) -> int:
        """
        Cuenta facturas emitidas en el mes que contienen el cupón en invoice.discounts[].source.coupon.

        Stripe v15+: los cupones de uso único se eliminan del cliente/suscripción tras el primer
        cobro, por lo que customer.discount queda vacío. La única fuente fiable es la factura
        generada en ese momento, que conserva el registro del descuento aplicado.
        Se deduplica por ID de suscripción para no contar más de una vez por suscripción.
        """
        _, last_day = calendar.monthrange(year, month)
        start = int(datetime(year, month, 1,        0,  0,  0, tzinfo=timezone.utc).timestamp())
        end   = int(datetime(year, month, last_day, 23, 59, 59, tzinfo=timezone.utc).timestamp())

        count     = 0
        seen_subs: set = set()
        has_more  = True
        last_id: Optional[str] = None

        while has_more:
            kwargs: dict = {
                "created": {"gte": start, "lte": end},
                "limit":   100,
                "expand":  ["data.discounts"],
            }
            if last_id:
                kwargs["starting_after"] = last_id

            try:
                invoices = _stripe.Invoice.list(**kwargs)
            except Exception as e:
                logger.error(f"ComercialReporter: Stripe error para {coupon_id}: {e}")
                break

            for inv in invoices.data:
                discounts = getattr(inv, "discounts", []) or []
                for d in discounts:
                    source     = getattr(d, "source", None)
                    src_coupon = getattr(source, "coupon", None) if source else None
                    if src_coupon == coupon_id:
                        sub_id = (
                            getattr(d, "subscription", None)
                            or getattr(inv, "subscription", None)
                            or inv.id
                        )
                        if sub_id not in seen_subs:
                            seen_subs.add(sub_id)
                            count += 1
                        break  # solo contar una vez por factura

            has_more = invoices.has_more
            if invoices.data:
                last_id = invoices.data[-1].id

        return count

    # ── Discord: conteo de miembros ─────────────────────────────

    async def _get_discord_member_count(self) -> int:
        if not self.discord_token:
            logger.warning("ComercialReporter: sin bot_token, member_count=0")
            return 0
        try:
            url = (
                f"https://discord.com/api/v10/guilds/{self.discord_guild_id}"
                "?with_counts=true"
            )
            headers = {"Authorization": f"Bot {self.discord_token}"}
            async with aiohttp.ClientSession() as s:
                async with s.get(url, headers=headers,
                                 timeout=aiohttp.ClientTimeout(total=10)) as resp:
                    if resp.status == 200:
                        data = await resp.json()
                        return data.get("approximate_member_count", 0)
                    text = await resp.text()
                    logger.error(f"ComercialReporter: Discord API {resp.status}: {text[:100]}")
        except Exception as e:
            logger.error(f"ComercialReporter: error obteniendo miembros Discord: {e}")
        return 0

    # ── Email Comercial (no colaborador) ────────────────────────

    async def _send_comercial_email(
        self, persona: dict, uses: int, year: int, month: int
    ) -> None:
        nombre  = persona["nombre"]
        mail    = persona["mail"]
        codigo  = persona["codigo"]
        mes_str = MESES_ES[month]
        subject = f"RESUMEN MES DE {mes_str} CANJES CODIGO {codigo}"

        if uses == 0:
            cuerpo_html = f"""
            <p>Buenos días {nombre},</p>
            <p>Otro mes que ha pasado, y otro mes haciendo balance de los canjes que puedas
            tener para disfrutar de la gratificación de 10€:</p>

            <table border="0" cellpadding="6"
                   style="border-collapse:collapse;font-family:Arial,sans-serif;">
              <tr><td style="color:#555;">Código</td>
                  <td><strong>{codigo}</strong></td></tr>
              <tr><td style="color:#555;">Tipo</td>
                  <td>Importe fijo de descuento</td></tr>
              <tr><td style="color:#555;">Condiciones</td>
                  <td>Descuento de 19,99 € una vez</td></tr>
              <tr><td style="color:#555;">Uso en {mes_str.capitalize()}</td>
                  <td>No hay canjes todavía</td></tr>
            </table>

            <p>Por desgracia todavía no ha habido ningún canje…
            vamos a ver el próximo mes! 😊</p>
            <p>Cualquier cosa y dudas, lo comentamos.</p>
            <p>Un saludo,</p>
            """ + FIRMA_HTML
        else:
            importe = uses * 10
            cuerpo_html = f"""
            <p>Buenos días {nombre},</p>
            <p>¡Enhorabuena! Este mes ha habido
            <strong>{uses} usuario{"s" if uses > 1 else ""}</strong>
            que se {"han" if uses > 1 else "ha"} dado de alta con tu código
            <strong>{codigo}</strong>, por lo que has devengado
            <strong>{importe}€</strong> a tu favor.</p>

            <table border="0" cellpadding="6"
                   style="border-collapse:collapse;font-family:Arial,sans-serif;">
              <tr><td style="color:#555;">Código</td>
                  <td><strong>{codigo}</strong></td></tr>
              <tr><td style="color:#555;">Canjes en {mes_str.capitalize()}</td>
                  <td><strong>{uses}</strong></td></tr>
              <tr><td style="color:#555;">Importe devengado</td>
                  <td><strong>{importe}€</strong></td></tr>
            </table>

            <p>Por favor, pásame <strong>factura o nota de gasto</strong> en donde
            aparezcan tus datos y número de cuenta bancaria para hacerte el ingreso.</p>
            <p>Cualquier cosa y dudas, lo comentamos.</p>
            <p>Un saludo,</p>
            """ + FIRMA_HTML

        self._send_email(mail, subject, cuerpo_html)
        logger.info(f"ComercialReporter: email enviado a {nombre} ({mail}) — {uses} canjes")

    # ── Email Colaborador (liquidación) ─────────────────────────

    async def _send_colaborador_email(
        self,
        persona: dict,
        uses: int,
        member_count: int,
        year: int,
        month: int,
        now: datetime,
    ) -> None:
        nombre  = persona["nombre"]
        mail    = persona["mail"]
        mes_str = MESES_ES[month]

        total_base = max(member_count - self.gratuitos, 0)
        honorarios = (
            total_base * 2.0
            if total_base <= 200
            else 200 * 2.0 + (total_base - 200) * 1.5
        )
        altas_codigo   = uses * 10.0
        total_base_fin = honorarios + altas_codigo
        a_pagar        = round(total_base_fin * 1.21, 2)
        fecha_hoy      = f"{now.day:02d}/{now.month:02d}/{now.year}"

        def fmt(v: float) -> str:
            return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        subject = f"LIQUIDACIÓN MES DE {mes_str}"

        cuerpo_html = f"""
        <p>Buenos días {nombre},</p>
        <p>Espero que estés bien.</p>
        <p>Estoy liquidando la parte de tus honorarios por colaborar con nosotros.</p>
        <p>Te paso un resumen de los números, a la espera de que me hagas llegar la
        factura y te lo pago al instante:</p>

        <table border="1" cellpadding="8" cellspacing="0"
               style="border-collapse:collapse;font-family:Arial,sans-serif;
                      font-size:14px;min-width:320px;">
          <tr>
            <td style="background:#f5f5f5;">Total miembros {fecha_hoy}</td>
            <td><strong>{member_count}</strong></td>
          </tr>
          <tr>
            <td style="background:#f5f5f5;">Gratuitos</td>
            <td>{self.gratuitos}</td>
          </tr>
          <tr><td colspan="2" style="padding:4px;"></td></tr>
          <tr>
            <td style="background:#f5f5f5;">Total base</td>
            <td>{total_base}</td>
          </tr>
          <tr>
            <td style="background:#f5f5f5;">Honorarios 2€</td>
            <td>{fmt(honorarios)}</td>
          </tr>
          <tr><td colspan="2" style="padding:4px;"></td></tr>
          <tr>
            <td style="background:#f5f5f5;">Masterclass</td>
            <td>-</td>
          </tr>
          <tr>
            <td style="background:#f5f5f5;">Altas el día Masterclass</td>
            <td>-</td>
          </tr>
          <tr>
            <td style="background:#f5f5f5;">Altas Código Embajador</td>
            <td>{fmt(altas_codigo)}</td>
          </tr>
          <tr><td colspan="2" style="padding:4px;"></td></tr>
          <tr>
            <td style="background:#e8e8e8;font-weight:bold;">TOTAL BASE</td>
            <td><strong>{fmt(total_base_fin)}</strong></td>
          </tr>
          <tr>
            <td style="background:#e8e8e8;font-weight:bold;">A pagar (IVA incluido)</td>
            <td><strong>{fmt(a_pagar)}</strong></td>
          </tr>
        </table>

        <p>Hazme llegar la factura y lo liquido al momento.</p>
        <p>Muchas gracias,</p>
        """ + FIRMA_HTML

        self._send_email(mail, subject, cuerpo_html)
        logger.info(
            f"ComercialReporter: liquidación enviada a {nombre} ({mail}) — "
            f"miembros={member_count} canjes={uses} a_pagar={a_pagar}€"
        )

    # ── Log a Discord ───────────────────────────────────────────

    async def _log_to_discord(
        self, results: list, mes_str: str, year: int, member_count: int
    ) -> None:
        if not self.log_webhook:
            return

        ok_count = sum(1 for r in results if r["ok"])
        all_ok   = ok_count == len(results)
        some_ok  = ok_count > 0
        color    = 0x2ECC71 if all_ok else (0xF39C12 if some_ok else 0xE74C3C)
        icon     = "✅" if all_ok else ("⚠️" if some_ok else "❌")

        fields = []
        for r in results:
            if r["ok"]:
                uses_txt = f"{r['uses']} canje{'s' if r['uses'] != 1 else ''}"
                value = f"✅ Email enviado · {uses_txt}"
            else:
                value = f"❌ Error: {r['error'] or 'desconocido'}"
            fields.append({
                "name":   f"{r['nombre']} — {r['tipo']}",
                "value":  f"`{r['mail']}`\n{value}",
                "inline": False,
            })

        fields.append({
            "name":   "📊 Miembros Discord",
            "value":  f"`{member_count}` miembros (fecha de envío)",
            "inline": True,
        })
        fields.append({
            "name":   "📬 Resultado global",
            "value":  f"`{ok_count}/{len(results)}` emails enviados correctamente",
            "inline": True,
        })

        embed = {
            "title":       f"{icon} Reporte Comerciales — {mes_str} {year}",
            "description": "Envío automático mensual a comerciales y colaboradores.",
            "color":       color,
            "fields":      fields,
            "timestamp":   datetime.utcnow().isoformat() + "Z",
            "footer":      {"text": "MTO Bot · comerciales_reporter"},
        }

        try:
            async with aiohttp.ClientSession() as s:
                async with s.post(
                    self.log_webhook,
                    json={"embeds": [embed]},
                    timeout=aiohttp.ClientTimeout(total=10),
                ) as resp:
                    if resp.status not in (200, 204):
                        text = await resp.text()
                        logger.warning(
                            f"ComercialReporter: log Discord HTTP {resp.status}: {text[:100]}"
                        )
                    else:
                        logger.info("ComercialReporter: log publicado en Discord ✅")
        except Exception as e:
            logger.error(f"ComercialReporter: error publicando log en Discord: {e}")

    # ── SMTP ────────────────────────────────────────────────────

    def _send_email(self, to: str, subject: str, html_body: str) -> None:
        msg = _build_email(
            from_addr = self.smtp_user,
            to_addr   = to,
            subject   = subject,
            html_body = html_body,
            cc_addr   = self.smtp_user,
        )

        ctx = ssl.create_default_context()
        with smtplib.SMTP(self.smtp_host, self.smtp_port) as server:
            server.ehlo()
            server.starttls(context=ctx)
            server.login(self.smtp_user, self.smtp_pass)
            server.sendmail(self.smtp_user, [to, self.smtp_user], msg.as_bytes())
        logger.info(f"ComercialReporter: email OK → {to} | {subject}")


def build_from_config(cfg: dict) -> Optional["ComercialReporter"]:
    com_cfg = cfg.get("comerciales", {})
    if not com_cfg.get("enabled", False):
        return None
    return ComercialReporter(cfg)
