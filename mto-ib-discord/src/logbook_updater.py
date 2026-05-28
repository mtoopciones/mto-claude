"""
Actualización automática del Log_book de Excel vía Dropbox API.

Flujo:
  1. Descarga el Excel desde Dropbox
  2. Modifica la hoja Log_book con openpyxl
  3. Sube de vuelta a Dropbox
  4. Mantiene un índice local JSON para relacionar aperturas con cierres

Columnas que escribe el bot (el resto son fórmulas Excel y NO se tocan):
  APERTURA: C, D, E, F, G, H, I, J, K, L, M, N, O, Q
  CIERRE:   T, U, V, W, X, Y, Z, AB
  ROLL:     AF  (mismo número en cierre y apertura del roll)
"""

import io
import json
import os
import re
import asyncio
import zipfile
from copy import copy
from datetime import datetime
from typing import Optional, Dict, Tuple
from xml.etree import ElementTree as ET
from zoneinfo import ZoneInfo

import aiohttp
from loguru import logger

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False
    logger.warning("openpyxl no disponible — logbook desactivado")

_MADRID = ZoneInfo("Europe/Madrid")


def _normalize_right(right: str) -> str:
    """Convierte P/C (formato IB) → PUT/CALL (columna Excel J/U)."""
    r = (right or "").upper()
    if r in ("P", "PUT"):
        return "PUT"
    if r in ("C", "CALL"):
        return "CALL"
    return right  # fallback sin cambios


def _norm_right_simple(r: str) -> str:
    """Normaliza P/PUT → 'P', C/CALL → 'C' para comparaciones."""
    r = (r or "").upper().strip()
    if r in ("P", "PUT"):
        return "P"
    if r in ("C", "CALL"):
        return "C"
    return r


def _restore_rich_data(original_bytes: bytes, new_bytes: bytes, sheet_name: str) -> bytes:
    """
    openpyxl elimina los atributos ``vm=`` de las celdas y los archivos
    ``xl/richData/*`` / ``xl/metadata.xml`` al guardar.
    Esta función los restaura desde el Excel original para que los tickers
    con tipo de dato 'Stock' no muestren #¡VALOR!.
    """
    import re as _re
    try:
        # ── 1. Leer del original: vm= por celda, archivos richData/metadata ──
        vm_map:     dict = {}   # {ref_str: vm_value_str}
        rich_files: dict = {}   # {path: bytes}
        orig_wb_rels_bytes = None
        orig_ct_bytes      = None
        orig_sheet_file    = None

        with zipfile.ZipFile(io.BytesIO(original_bytes)) as oz:
            znames = set(oz.namelist())

            for n in znames:
                if n.startswith("xl/richData/") or n == "xl/metadata.xml":
                    rich_files[n] = oz.read(n)

            if not rich_files:
                return new_bytes  # sin Stock data type → nada que restaurar

            # Localizar archivo de hoja en el original
            wb_root = ET.fromstring(oz.open("xl/workbook.xml").read())
            rel_id_orig = None
            for s in wb_root.iter():
                if s.tag.split("}")[-1] == "sheet" and s.get("name") == sheet_name:
                    rel_id_orig = s.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )
                    break
            if rel_id_orig and "xl/_rels/workbook.xml.rels" in znames:
                rels = ET.fromstring(oz.open("xl/_rels/workbook.xml.rels").read())
                for r in rels.iter():
                    if r.tag.split("}")[-1] == "Relationship" and r.get("Id") == rel_id_orig:
                        tgt = r.get("Target", "").lstrip("/")
                        orig_sheet_file = tgt if tgt.startswith("xl/") else f"xl/{tgt}"
                        break

            if orig_sheet_file and orig_sheet_file in znames:
                for c in ET.fromstring(oz.open(orig_sheet_file).read()).iter():
                    if c.tag.split("}")[-1] == "c":
                        vm  = c.get("vm")
                        ref = c.get("r")
                        if vm and ref:
                            vm_map[ref] = vm

            if "xl/_rels/workbook.xml.rels" in znames:
                orig_wb_rels_bytes = oz.read("xl/_rels/workbook.xml.rels")
            if "[Content_Types].xml" in znames:
                orig_ct_bytes = oz.read("[Content_Types].xml")

        if not vm_map:
            return new_bytes  # ninguna celda con vm= → nada que restaurar

        # ── 2. Localizar hoja en el nuevo zip ──────────────────────────────
        new_sheet_file = None
        with zipfile.ZipFile(io.BytesIO(new_bytes)) as nz:
            wb_root = ET.fromstring(nz.open("xl/workbook.xml").read())
            rel_id_new = None
            for s in wb_root.iter():
                if s.tag.split("}")[-1] == "sheet" and s.get("name") == sheet_name:
                    rel_id_new = s.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )
                    break
            if rel_id_new and "xl/_rels/workbook.xml.rels" in set(nz.namelist()):
                rels = ET.fromstring(nz.open("xl/_rels/workbook.xml.rels").read())
                for r in rels.iter():
                    if r.tag.split("}")[-1] == "Relationship" and r.get("Id") == rel_id_new:
                        tgt = r.get("Target", "").lstrip("/")
                        new_sheet_file = tgt if tgt.startswith("xl/") else f"xl/{tgt}"
                        break

        # ── 3. Reconstruir zip con richData + metadata + vm= restaurados ───
        in_buf  = io.BytesIO(new_bytes)
        out_buf = io.BytesIO()

        with zipfile.ZipFile(in_buf, "r") as old_z, \
             zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as new_z:

            existing_names = set(old_z.namelist())

            for item in old_z.namelist():
                data = old_z.read(item)

                if item == new_sheet_file and vm_map:
                    # Restaurar atributo vm= en celdas que lo tenían
                    xml_str = data.decode("utf-8")

                    def _repl(m, _vmap=vm_map):
                        tag = m.group(0)
                        rm  = _re.search(r'\br="([^"]+)"', tag)
                        if rm:
                            ref = rm.group(1)
                            if ref in _vmap and 'vm="' not in tag:
                                tag = tag.replace(
                                    f'r="{ref}"', f'r="{ref}" vm="{_vmap[ref]}"', 1
                                )
                        return tag

                    xml_str = _re.sub(r'<c\b[^>]*>', _repl, xml_str)
                    data = xml_str.encode("utf-8")

                elif item in rich_files:
                    data = rich_files[item]  # restaurar desde original

                elif item == "[Content_Types].xml" and orig_ct_bytes:
                    orig_ct = ET.fromstring(orig_ct_bytes)
                    new_ct  = ET.fromstring(data)
                    exist_parts = {c.get("PartName", "") for c in new_ct}
                    for child in orig_ct:
                        pn = child.get("PartName", "")
                        ct = child.get("ContentType", "")
                        if pn and pn not in exist_parts and (
                            "richData" in pn or "richData" in ct
                            or "metadata" in pn.lower()
                        ):
                            new_ct.append(child)
                            exist_parts.add(pn)
                    data = (
                        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        + ET.tostring(new_ct, encoding="unicode").encode("utf-8")
                    )

                elif item == "xl/_rels/workbook.xml.rels" and orig_wb_rels_bytes:
                    orig_rels = ET.fromstring(orig_wb_rels_bytes)
                    new_rels  = ET.fromstring(data)
                    exist_tgts = {r.get("Target", "") for r in new_rels}
                    exist_ids  = {r.get("Id",     "") for r in new_rels}
                    max_id = max(
                        (int(i[3:]) for i in exist_ids
                         if i.startswith("rId") and i[3:].isdigit()),
                        default=0,
                    )
                    for r in orig_rels:
                        tgt = r.get("Target", "")
                        if tgt and tgt not in exist_tgts and (
                            "richData" in tgt or "metadata" in tgt.lower()
                        ):
                            max_id += 1
                            new_r = ET.SubElement(new_rels, r.tag)
                            new_r.set("Id",     f"rId{max_id}")
                            new_r.set("Type",   r.get("Type", ""))
                            new_r.set("Target", tgt)
                            exist_tgts.add(tgt)
                    data = (
                        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        + ET.tostring(new_rels, encoding="unicode").encode("utf-8")
                    )

                new_z.writestr(item, data)

            # Agregar archivos richData que no existían en el nuevo zip
            for fname, fdata in rich_files.items():
                if fname not in existing_names:
                    new_z.writestr(fname, fdata)

        return out_buf.getvalue()

    except Exception as e:
        logger.warning(f"_restore_rich_data: {e} — subiendo sin restaurar richData")
        return new_bytes  # fallar de forma segura


def _extract_rich_cell_values(xlsx_bytes: bytes, sheet_name: str) -> dict:
    """
    Lee los valores de celdas con tipo de dato 'Stock' de Excel (rich values).
    Excel guarda la info en xl/richData/ — openpyxl no lo soporta.
    Devuelve {(row, col): "Nombre empresa (EXCHANGE:TICKER)"}.

    Cadena de índices:
        celda.vm (1-based)
          → valueMetadata.bk[vm-1].rc.v           = fm_idx
          → futureMetadata[XLRICHVALUE].bk[fm_idx]
              .extLst…rvb.i                        = rv_idx   (puntero _linkedentity)
          → rv_vals[rv_idx].v[0]                  = core_rv_idx  (si es _linkedentity)
          → rv_vals[core_rv_idx].v[7]             = _DisplayString
    """
    # Textos de UI de Microsoft que NO son nombre de empresa
    _UI_LABELS = {
        "learn more on bing", "learn more", "bing",
        "learn more about this data type",
    }

    result: dict = {}
    try:
        with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as z:
            znames = set(z.namelist())

            # ── 1. Rich values: rv_vals[i] = (s_type_int, [v0, v1, …]) ────
            rv_vals: list = []
            for rdf in ("xl/richData/rdrichvalue.xml", "xl/richData/rdRichValue.xml"):
                if rdf not in znames:
                    continue
                root = ET.fromstring(z.open(rdf).read())
                for rv in root.iter():
                    if rv.tag.split("}")[-1] == "rv":
                        rv_vals.append((
                            int(rv.get("s", 0)),
                            [ch.text or "" for ch in rv if ch.tag.split("}")[-1] == "v"],
                        ))
                break
            if not rv_vals:
                return result

            # ── 2. Schema: rvtype_ds_idx[s_type] = índice de _DisplayString ─
            # _DisplayString es el texto que Excel muestra en la celda.
            # En todos los _linkedentitycore aparece en posición 7.
            rvtype_ds_idx: dict = {}
            rvs_path = "xl/richData/rdrichvaluestructure.xml"
            if rvs_path in znames:
                rvsroot = ET.fromstring(z.open(rvs_path).read())
                for ti, typ in enumerate(rvsroot):
                    if typ.tag.split("}")[-1] not in ("rvTyp", "rvType", "s"):
                        continue
                    keys = [k for k in typ if k.tag.split("}")[-1] == "k"]
                    for i, k in enumerate(keys):
                        if k.get("n") == "_DisplayString":
                            rvtype_ds_idx[ti] = i
                            break

            # ── 3. metadata.xml ─────────────────────────────────────────────
            # futureMetadata[XLRICHVALUE]:  fm_idx → rv_idx  (rvb.i, dentro de extLst)
            # valueMetadata:                vm_idx → fm_idx  (rc.v)
            fm_to_rv: list = []   # índice = fm_bk_idx, valor = rv_idx
            vm_to_fm: list = []   # índice = vm_bk_idx, valor = fm_idx

            if "xl/metadata.xml" in znames:
                meta = ET.fromstring(z.open("xl/metadata.xml").read())
                for sect in meta:
                    tag = sect.tag.split("}")[-1]

                    # Solo procesar la sección XLRICHVALUE (ignorar XLDAPR, etc.)
                    if tag == "futureMetadata" and sect.get("name") == "XLRICHVALUE":
                        for bk in sect:
                            if bk.tag.split("}")[-1] != "bk":
                                continue
                            rv_i = -1
                            for elem in bk.iter():   # rvb está dentro de extLst
                                if elem.tag.split("}")[-1] == "rvb":
                                    try:
                                        rv_i = int(elem.get("i", -1))
                                    except Exception:
                                        pass
                                    break
                            fm_to_rv.append(rv_i)

                    elif tag == "valueMetadata":
                        for bk in sect:
                            if bk.tag.split("}")[-1] != "bk":
                                continue
                            fm_j = -1
                            for rc in bk:            # rc es hijo directo de bk
                                if rc.tag.split("}")[-1] == "rc":
                                    try:
                                        fm_j = int(rc.get("v", -1))
                                    except Exception:
                                        pass
                                    break
                            vm_to_fm.append(fm_j)

            # ── 4. Localizar el archivo de la hoja ─────────────────────────
            wb_root = ET.fromstring(z.open("xl/workbook.xml").read())
            rel_id = None
            for s in wb_root.iter():
                if s.tag.split("}")[-1] == "sheet" and s.get("name") == sheet_name:
                    rel_id = s.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )
                    break
            if not rel_id:
                return result

            rels = ET.fromstring(z.open("xl/_rels/workbook.xml.rels").read())
            sheet_file = None
            for r in rels.iter():
                if r.tag.split("}")[-1] == "Relationship" and r.get("Id") == rel_id:
                    tgt = r.get("Target", "").lstrip("/")
                    sheet_file = tgt if tgt.startswith("xl/") else f"xl/{tgt}"
                    break
            if not sheet_file or sheet_file not in znames:
                return result

            # ── 5. Parsear celdas con atributo vm (= Stock data type) ──────
            if not OPENPYXL_OK:
                return result
            from openpyxl.utils.cell import coordinate_to_tuple

            ws_xml = ET.fromstring(z.open(sheet_file).read())
            for c in ws_xml.iter():
                if c.tag.split("}")[-1] != "c":
                    continue
                vm = c.get("vm")
                if vm is None:
                    continue
                try:
                    row, col = coordinate_to_tuple(c.get("r", ""))
                except Exception:
                    continue

                # vm (1-based) → fm_idx → rv_idx
                vm_0   = int(vm) - 1
                fm_idx = vm_to_fm[vm_0] if 0 <= vm_0 < len(vm_to_fm) else -1
                if fm_idx < 0:
                    continue
                rv_idx = fm_to_rv[fm_idx] if 0 <= fm_idx < len(fm_to_rv) else -1
                if rv_idx < 0 or rv_idx >= len(rv_vals):
                    continue

                s_type, vals = rv_vals[rv_idx]

                # Si es _linkedentity (1 prop = índice del _linkedentitycore): seguir puntero
                if len(vals) == 1 and vals[0].isdigit():
                    core_idx = int(vals[0])
                    if 0 <= core_idx < len(rv_vals):
                        s_type, vals = rv_vals[core_idx]

                if not vals:
                    continue

                # _DisplayString: índice 7 en todos los _linkedentitycore (siempre constante)
                ds_idx  = rvtype_ds_idx.get(s_type, 7)
                display = vals[ds_idx] if ds_idx < len(vals) else ""

                # Fallback: primer texto no-URL, no-numérico, no-UI, > 3 chars
                if not display or display.lower() in _UI_LABELS:
                    display = ""
                    for v in vals:
                        if (v
                                and not v.startswith("http")
                                and v.lower() not in _UI_LABELS
                                and not v.replace(".", "").replace(",", "").replace("-", "").isnumeric()
                                and len(v) > 3):
                            display = v
                            break

                if display:
                    result[(row, col)] = display

    except Exception as e:
        logger.debug(f"Rich cell values extraction: {e}")
    return result


# ── Mapeo de nombres de cuenta → Canal en el Excel ────────────
CANAL_MAP = {
    "MTO Cuenta 10K":        "10 k",
    "MTO Cuenta 50k":        "50 k",
    "MTO Dividendos ETF 5K": "Dividendos - ETF",
}

# ── Columnas del Log_book (1-based, openpyxl) ─────────────────
COL = {
    "B":  2,   # Num. Registro (pre-llenado)
    "C":  3,   # VALOR (ticker)
    "D":  4,   # Estrategia
    "E":  5,   # Canal
    "F":  6,   # COTIZACIÓN
    "G":  7,   # STRIKE
    "H":  8,   # FECHA AP.
    "I":  9,   # VENCIMIENTO
    "J":  10,  # TIPO (PUT/CALL)
    "K":  11,  # OPERACIÓN (Venta/Compra)
    "L":  12,  # LOTE
    "M":  13,  # CANTIDAD
    "N":  14,  # PRIMA
    "O":  15,  # COMISIÓN
    # P=16 es FÓRMULA → no tocar
    "Q":  17,  # ESTADO apertura
    # R=18 separador
    # S=19 es FÓRMULA (OTM/ITM) → no tocar
    "T":  20,  # FECHA CIERRE
    "U":  21,  # TIPO cierre (PUT/CALL)
    "V":  22,  # OPERACIÓN cierre (Venta/Compra)
    "W":  23,  # LOTE cierre
    "X":  24,  # CANTIDAD cierre
    "Y":  25,  # PRIMA cierre
    "Z":  26,  # COMISIÓN cierre
    # AA=27 es FÓRMULA → no tocar
    "AB": 28,  # ESTADO cierre
    # AC=29 separador
    # AD=30 FÓRMULA P&L
    # AE=31 FÓRMULA Cash en curso
    "AF": 32,  # Ref. Roll (manual)
    "AG": 33,  # FÓRMULA Fecha cómputo (solo se copia el formato/fórmula, no se escribe valor)
}

# Fila donde empieza el primer registro de datos
DATA_START_ROW = 4
SHEET_NAME     = "Log_book"


class LogbookUpdater:
    def __init__(
        self,
        dropbox_app_key:      str,
        dropbox_app_secret:   str,
        dropbox_refresh_token: str,
        dropbox_path:         str,
        index_file:           str = "data/logbook_index.json",
    ):
        """
        dropbox_app_key/secret   : credenciales de la app Dropbox
        dropbox_refresh_token    : refresh token permanente (no caduca)
        dropbox_path             : ruta del Excel en Dropbox
        index_file               : JSON local que mapea apertura → fila Excel
        """
        self.app_key       = dropbox_app_key
        self.app_secret    = dropbox_app_secret
        self.refresh_token = dropbox_refresh_token
        self.dropbox_path  = dropbox_path
        self.index_file    = index_file
        self.log_channel   = None
        self._lock         = asyncio.Lock()
        self._access_token: Optional[str] = None   # se renueva automáticamente
        self._root_namespace_id: Optional[str] = None   # team namespace (Business)
        self._index: Dict[str, int] = {}
        self._load_index()

    # ── Índice local ──────────────────────────────────────────

    def _load_index(self) -> None:
        try:
            if os.path.exists(self.index_file):
                with open(self.index_file, "r", encoding="utf-8") as f:
                    self._index = json.load(f)
                logger.info(f"Logbook index: {len(self._index)} entradas cargadas")
        except Exception as e:
            logger.error(f"Logbook: error cargando índice: {e}")
            self._index = {}

    def _save_index(self) -> None:
        try:
            os.makedirs(os.path.dirname(self.index_file), exist_ok=True)
            with open(self.index_file, "w", encoding="utf-8") as f:
                json.dump(self._index, f, indent=2, ensure_ascii=False)
        except Exception as e:
            logger.error(f"Logbook: error guardando índice: {e}")

    @staticmethod
    def _index_key(symbol: str, strike: float, expiry: str,
                   right: str, canal: str) -> str:
        return f"{symbol}|{strike}|{expiry}|{right}|{canal}"

    # ── Dropbox auth ──────────────────────────────────────────

    async def _get_access_token(self) -> Optional[str]:
        """Obtiene un access token fresco usando el refresh token."""
        try:
            data = {
                "grant_type":    "refresh_token",
                "refresh_token": self.refresh_token,
                "client_id":     self.app_key,
                "client_secret": self.app_secret,
            }
            async with aiohttp.ClientSession() as s:
                async with s.post(
                    "https://api.dropboxapi.com/oauth2/token",
                    data=data
                ) as resp:
                    if resp.status == 200:
                        result = await resp.json()
                        return result.get("access_token")
                    text = await resp.text()
                    logger.error(f"Dropbox token refresh error {resp.status}: {text[:200]}")
                    return None
        except Exception as e:
            logger.error(f"Dropbox token refresh exception: {e}")
            return None

    # ── Dropbox helpers ───────────────────────────────────────

    async def _get_path_root_header(self, token: str) -> dict:
        """Devuelve el header Dropbox-API-Path-Root para acceder al team namespace."""
        if not self._root_namespace_id:
            try:
                async with aiohttp.ClientSession() as s:
                    async with s.post(
                        "https://api.dropboxapi.com/2/users/get_current_account",
                        headers={"Authorization": f"Bearer {token}",
                                 "Content-Type": "application/json"},
                        data=b"null"
                    ) as resp:
                        if resp.status == 200:
                            acc = await resp.json()
                            self._root_namespace_id = (
                                acc.get("root_info", {}).get("root_namespace_id")
                            )
            except Exception as e:
                logger.warning(f"Dropbox: no se pudo obtener root namespace: {e}")
        if self._root_namespace_id:
            return {"Dropbox-API-Path-Root": json.dumps({
                ".tag": "namespace_id",
                "namespace_id": self._root_namespace_id,
            })}
        return {}

    async def _download(self) -> Optional[bytes]:
        token = await self._get_access_token()
        if not token:
            return None
        path_root = await self._get_path_root_header(token)
        headers = {
            "Authorization":   f"Bearer {token}",
            "Dropbox-API-Arg": json.dumps({"path": self.dropbox_path}),
            **path_root,
        }
        async with aiohttp.ClientSession() as s:
            async with s.post(
                "https://content.dropboxapi.com/2/files/download",
                headers=headers
            ) as resp:
                if resp.status == 200:
                    return await resp.read()
                text = await resp.text()
                logger.error(f"Dropbox download error {resp.status}: {text[:200]}")
                return None

    async def _upload(self, data: bytes) -> bool:
        token = await self._get_access_token()
        if not token:
            return False
        path_root = await self._get_path_root_header(token)
        headers = {
            "Authorization":   f"Bearer {token}",
            "Dropbox-API-Arg": json.dumps({
                "path":       self.dropbox_path,
                "mode":       "overwrite",
                "autorename": False,
                "mute":       True,
            }),
            "Content-Type": "application/octet-stream",
            **path_root,
        }
        async with aiohttp.ClientSession() as s:
            async with s.post(
                "https://content.dropboxapi.com/2/files/upload",
                headers=headers,
                data=data
            ) as resp:
                if resp.status == 200:
                    return True
                text = await resp.text()
                logger.error(f"Dropbox upload error {resp.status}: {text[:200]}")
                return False

    # ── Operaciones sobre el workbook ─────────────────────────

    def _find_free_row(self, ws) -> int:
        """Devuelve la primera fila con B pre-llenado pero D vacío."""
        for row in range(DATA_START_ROW, ws.max_row + 50):
            b = ws.cell(row=row, column=COL["B"]).value
            d = ws.cell(row=row, column=COL["D"]).value
            if b is not None and d is None:
                return row
        # Si no hay filas pre-llenadas, usar la siguiente al último dato
        for row in range(ws.max_row, DATA_START_ROW - 1, -1):
            if ws.cell(row=row, column=COL["B"]).value is not None:
                return row + 1
        return DATA_START_ROW

    def _next_roll_number(self, ws) -> int:
        """Devuelve el siguiente número de roll libre (máx actual + 1)."""
        max_roll = 0
        for row in range(DATA_START_ROW, ws.max_row + 1):
            val = ws.cell(row=row, column=COL["AF"]).value
            if isinstance(val, (int, float)) and val > max_roll:
                max_roll = int(val)
        return max_roll + 1

    @staticmethod
    def _naive(dt):
        """Elimina tzinfo de un datetime para compatibilidad con openpyxl/Excel."""
        if isinstance(dt, datetime) and dt.tzinfo is not None:
            return dt.replace(tzinfo=None)
        return dt

    def _find_style_source_row(self, ws, target_row: int) -> Optional[int]:
        """Devuelve la fila anterior más cercana que tenga fecha de apertura (col H rellena)."""
        for r in range(target_row - 1, DATA_START_ROW - 1, -1):
            if ws.cell(row=r, column=COL["H"]).value is not None:
                return r
        return None

    def _copy_row_style(self, ws, src_row: int, dst_row: int) -> None:
        """
        Copia estilos (number_format, font, fill, border, alignment) y fórmulas
        de src_row a dst_row, ajustando automáticamente los números de fila en
        las fórmulas con referencias relativas (no afecta referencias absolutas $).

        Garantiza:
          - Fechas con formato DD/MM/YYYY en columnas H, I, T
          - Cotización en formato numérico/precio en columna F
          - Fórmulas arrastradas en columnas P, S, AA, AD, AE, AG
          - Todos los formatos de número, fuente y borde coherentes
        """
        delta    = dst_row - src_row
        max_col  = max(ws.max_column or 0, COL["AG"])

        for col in range(1, max_col + 1):
            src_cell = ws.cell(row=src_row, column=col)
            dst_cell = ws.cell(row=dst_row, column=col)

            # ── Estilos ────────────────────────────────────────────────────
            if src_cell.has_style:
                dst_cell.font          = copy(src_cell.font)
                dst_cell.fill          = copy(src_cell.fill)
                dst_cell.border        = copy(src_cell.border)
                dst_cell.alignment     = copy(src_cell.alignment)
                dst_cell.number_format = src_cell.number_format

            # ── Fórmulas: copiar y ajustar referencias relativas de fila ──
            v = src_cell.value
            if isinstance(v, str) and v.startswith("="):
                # Ajusta A1 y AA1 (relativas) pero NO $A1, A$1 ni $A$1 (absolutas)
                adjusted = re.sub(
                    r'(?<!\$)([A-Z]+)(?<!\$)(\d+)',
                    lambda m: m.group(1) + str(int(m.group(2)) + delta),
                    v,
                )
                dst_cell.value = adjusted

    def _write_open(self, ws, row: int, data: dict) -> None:
        """Escribe los campos de apertura en la fila indicada."""
        # ── Paso 1: copiar estilos y fórmulas de la fila anterior ─────────
        # Esto garantiza formatos de fecha (DD/MM/YYYY), formato de cotización,
        # y arrastra las fórmulas de Excel (columnas P, S, AA, AD, AE, AG).
        src_row = self._find_style_source_row(ws, row)
        if src_row:
            self._copy_row_style(ws, src_row, row)

        # ── Paso 2: escribir valores (sobreescriben los de _copy_row_style) ─
        ws.cell(row=row, column=COL["C"]).value = data["symbol"]
        ws.cell(row=row, column=COL["D"]).value = data["strategy"]
        ws.cell(row=row, column=COL["E"]).value = data["canal"]
        ws.cell(row=row, column=COL["F"]).value = data.get("stock_price")
        ws.cell(row=row, column=COL["G"]).value = data["strike"]
        ws.cell(row=row, column=COL["H"]).value = self._naive(data["trade_date"])
        ws.cell(row=row, column=COL["I"]).value = self._naive(data["expiry"])
        ws.cell(row=row, column=COL["J"]).value = data["right"]         # PUT / CALL
        ws.cell(row=row, column=COL["K"]).value = data["action"]        # Venta / Compra
        ws.cell(row=row, column=COL["L"]).value = 1                     # LOTE siempre 1
        ws.cell(row=row, column=COL["M"]).value = data["quantity"]
        ws.cell(row=row, column=COL["N"]).value = data["premium"]
        ws.cell(row=row, column=COL["O"]).value = data["commission"]
        ws.cell(row=row, column=COL["Q"]).value = "Abierta"
        if data.get("roll_ref"):
            ws.cell(row=row, column=COL["AF"]).value = data["roll_ref"]

    def _write_close(self, ws, row: int, data: dict) -> None:
        """Rellena los campos de cierre y actualiza Q en la fila indicada."""
        # Copiar formato de fecha de cierre (col T) desde la fila anterior con cierre.
        # (Si _write_open ya fue llamado con el nuevo código, el formato ya estará
        # copiado; esto solo es necesario para filas abiertas por código antiguo.)
        for r in range(row - 1, DATA_START_ROW - 1, -1):
            src_t = ws.cell(row=r, column=COL["T"])
            if src_t.value is not None and src_t.has_style:
                dst_t = ws.cell(row=row, column=COL["T"])
                dst_t.number_format = src_t.number_format
                break

        ws.cell(row=row, column=COL["T"]).value  = self._naive(data["close_date"])
        ws.cell(row=row, column=COL["U"]).value  = data["right"]
        ws.cell(row=row, column=COL["V"]).value  = data["action"]       # Venta / Compra
        ws.cell(row=row, column=COL["W"]).value  = 1
        ws.cell(row=row, column=COL["X"]).value  = data["quantity"]
        ws.cell(row=row, column=COL["Y"]).value  = data["premium"]
        ws.cell(row=row, column=COL["Z"]).value  = data["commission"]
        ws.cell(row=row, column=COL["AB"]).value = data.get("close_status", "Vendida")
        ws.cell(row=row, column=COL["Q"]).value  = data.get("open_status", "Cerrada")
        if data.get("roll_ref"):
            ws.cell(row=row, column=COL["AF"]).value = data["roll_ref"]

    # ── API pública ───────────────────────────────────────────

    async def record_open(
        self,
        symbol:      str,
        strategy:    str,
        account_name: str,
        strike:      float,
        expiry:      datetime,
        right:       str,
        action:      str,
        quantity:    int,
        premium:     float,
        commission:  float,
        trade_date:  datetime,
        stock_price: Optional[float] = None,
        roll_ref:    Optional[int]   = None,
    ) -> None:
        if not OPENPYXL_OK:
            return
        canal = CANAL_MAP.get(account_name, account_name)
        data = {
            "symbol":      symbol,
            "strategy":    strategy,
            "canal":       canal,
            "stock_price": stock_price,
            "strike":      strike,
            "trade_date":  trade_date,
            "expiry":      expiry,
            "right":       _normalize_right(right),   # "PUT" / "CALL"
            "action":      "Venta" if action.upper() == "SELL" else "Compra",
            "quantity":    quantity,
            "premium":     premium,
            "commission":  commission,
            "roll_ref":    roll_ref,
        }
        await self._update_excel("open", data, symbol, strike,
                                 expiry, right, canal)

    async def record_close(
        self,
        symbol:       str,
        account_name: str,
        strike:       float,
        expiry:       datetime,
        right:        str,
        action:       str,
        quantity:     int,
        premium:      float,
        commission:   float,
        close_date:   datetime,
        close_status: str = "Vendida",
        open_status:  str = "Cerrada",
        roll_ref:     Optional[int] = None,
    ) -> None:
        if not OPENPYXL_OK:
            return
        canal = CANAL_MAP.get(account_name, account_name)
        data = {
            "close_date":   close_date,
            "right":        _normalize_right(right),   # "PUT" / "CALL"
            "action":       "Venta" if action.upper() == "SELL" else "Compra",
            "quantity":     quantity,
            "premium":      premium,
            "commission":   commission,
            "close_status": close_status,
            "open_status":  open_status,
            "roll_ref":     roll_ref,
        }
        await self._update_excel("close", data, symbol, strike,
                                 expiry, right, canal)

    # ── Motor principal ───────────────────────────────────────

    async def _update_excel(
        self,
        op_type: str,
        data:    dict,
        symbol:  str,
        strike:  float,
        expiry:  datetime,
        right:   str,
        canal:   str,
    ) -> None:
        async with self._lock:
            try:
                raw = await self._download()
                if raw is None:
                    raise RuntimeError("No se pudo descargar el Excel de Dropbox")

                wb = openpyxl.load_workbook(io.BytesIO(raw))
                ws = wb[SHEET_NAME]

                expiry_str = expiry.strftime("%Y-%m-%d") if expiry else ""
                key = self._index_key(symbol, strike, expiry_str, right.upper(), canal)

                if op_type == "open":
                    row = self._find_free_row(ws)
                    num_registro = ws.cell(row=row, column=COL["B"]).value
                    self._write_open(ws, row, data)
                    # Guardar en índice para poder cerrar después
                    self._index[key] = row
                    self._save_index()
                    desc = f"APERTURA | {symbol} | fila {row} (reg. {num_registro}) | {canal}"

                elif op_type == "close":
                    row = self._index.get(key)
                    if row is not None:
                        # Sanidad: verificar que la fila del índice está realmente abierta
                        # (columna T vacía). Si ya tiene cierre, el índice está desactualizado.
                        t_existing = ws.cell(row=row, column=COL["T"]).value
                        if t_existing is not None:
                            logger.warning(
                                f"Logbook: índice apunta a fila {row} pero ya tiene cierre "
                                f"— buscando la fila correcta"
                            )
                            row = None

                    if row is None:
                        # Fallback: buscar en el Excel por coincidencia de campos.
                        # Se extrae el mapa de celdas Stock (rich values) para poder
                        # leer la columna C (ticker) que openpyxl devuelve como None.
                        rich_values = _extract_rich_cell_values(raw, SHEET_NAME)
                        row = self._find_open_row(ws, symbol, strike, expiry,
                                                  right, canal, rich_values)
                    if row is None:
                        logger.warning(
                            f"Logbook: no se encontró apertura para {key} — "
                            "anotando el cierre en nueva fila"
                        )
                        row = self._find_free_row(ws)

                    self._write_close(ws, row, data)
                    # Eliminar del índice (ya está cerrada)
                    self._index.pop(key, None)
                    self._save_index()
                    num_registro = ws.cell(row=row, column=COL["B"]).value
                    desc = f"CIERRE | {symbol} | fila {row} (reg. {num_registro}) | {canal}"

                else:
                    return

                # Serializar, restaurar richData (vm=) y subir
                buf = io.BytesIO()
                wb.save(buf)
                patched = _restore_rich_data(raw, buf.getvalue(), SHEET_NAME)
                ok = await self._upload(patched)

                if ok:
                    logger.info(f"Logbook actualizado: {desc}")
                    if self.log_channel:
                        await self.log_channel.send_info(
                            f"📋 Logbook actualizado y guardado en Excel: {desc}"
                        )
                else:
                    raise RuntimeError("Error al subir el Excel a Dropbox")

            except Exception as e:
                logger.error(f"Logbook error ({op_type}): {e}")
                if self.log_channel:
                    await self.log_channel.send_error(
                        f"No se pudo actualizar el logbook Excel ({op_type}): {e}"
                    )

    def _find_open_row(self, ws, symbol: str, strike: float,
                       expiry: datetime, right: str, canal: str,
                       rich_values: Optional[dict] = None) -> Optional[int]:
        """
        Busca en el Excel la fila de apertura por coincidencia de campos.

        rich_values: dict {(row, col): display_str} para leer celdas con tipo
                     de dato 'Stock' que openpyxl no puede leer directamente.
        """
        expiry_date = expiry.date() if expiry else None
        right_norm  = _norm_right_simple(right)   # "P" o "C"

        for row in range(DATA_START_ROW, ws.max_row + 1):
            c_sym = ws.cell(row=row, column=COL["C"]).value

            # Las celdas con tipo de dato 'Stock' devuelven None en openpyxl.
            # Usamos el mapa de rich values (XML) para obtener el ticker.
            if c_sym is None and rich_values:
                display = rich_values.get((row, COL["C"]), "")
                if display:
                    # Extraer ticker de "(EXCHANGE:TICKER)" al final del string
                    m = re.search(r'\(([^:)]+):([^)]+)\)\s*$', display)
                    c_sym = m.group(2).strip() if m else display.split()[0]

            if not c_sym or str(c_sym).upper() != symbol.upper():
                continue

            g_strike = ws.cell(row=row, column=COL["G"]).value
            i_expiry = ws.cell(row=row, column=COL["I"]).value
            j_right  = ws.cell(row=row, column=COL["J"]).value
            e_canal  = ws.cell(row=row, column=COL["E"]).value
            q_state  = ws.cell(row=row, column=COL["Q"]).value
            t_close  = ws.cell(row=row, column=COL["T"]).value

            # Strike: comparar con tolerancia para flotantes
            strike_ok = (g_strike is not None
                         and abs(float(g_strike) - (strike or 0)) < 0.01)

            if not (
                strike_ok
                and j_right and _norm_right_simple(j_right) == right_norm
                and (e_canal or "").strip() == canal.strip()
                and q_state in ("Abierta", None)
                and t_close is None
            ):
                continue

            if i_expiry:
                i_date = i_expiry.date() if hasattr(i_expiry, "date") else None
                if expiry_date and i_date and i_date == expiry_date:
                    return row
            else:
                return row
        return None

    # ── Ayuda para rolls ──────────────────────────────────────

    async def get_next_roll_ref(self) -> int:
        """Descarga el Excel y calcula el siguiente número de roll libre."""
        raw = await self._download()
        if raw is None:
            return 1
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[SHEET_NAME]
        return self._next_roll_number(ws)

    # ── Exportación semanal ───────────────────────────────────

    async def export_and_publish(self, discord_webhook: str) -> None:
        """
        Cada sábado a las 11:40 Madrid:
          1. Descarga el Excel maestro de Dropbox
          2. Extrae Log_book pegando TODO como valores (sin fórmulas)
          3. Sube el nuevo Excel a Dropbox en DOCUMENTOS PUBLICADOS EN DISCORD/
          4. Lo publica en Discord como adjunto
        """
        if not OPENPYXL_OK:
            logger.error("Logbook export: openpyxl no disponible")
            return

        fecha = datetime.now(_MADRID)
        filename = fecha.strftime("%Y.%m.%d") + " Excel operaciones.xlsx"
        dest_dropbox_path = (
            "/Easy Tax Advice/CLIENTES/74141 MTO OPCIONES, S.L"
            "/COMPARTIDA MTO/PUBLICACIONES/DOCUMENTOS PUBLICADOS EN DISCORD"
            f"/{filename}"
        )

        try:
            # 1. Descargar fuente (data_only=True para obtener valores calculados)
            raw = await self._download()
            if raw is None:
                raise RuntimeError("No se pudo descargar el Excel de Dropbox")

            wb_src = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            if SHEET_NAME not in wb_src.sheetnames:
                raise RuntimeError(f"Hoja '{SHEET_NAME}' no encontrada en el Excel")
            ws_src = wb_src[SHEET_NAME]

            # Extraer valores de celdas Stock (rich values) — openpyxl devuelve None en ellas
            rich_values = _extract_rich_cell_values(raw, SHEET_NAME)
            if rich_values:
                # Mostrar rango de filas donde se encontró VALOR (columna C = col 3)
                col_c_rows = sorted(r for (r, c) in rich_values if c == 3)
                row_range = f"filas {col_c_rows[0]}–{col_c_rows[-1]}" if col_c_rows else "ninguna fila en col C"
                logger.info(
                    f"Logbook export: {len(rich_values)} celdas Stock detectadas "
                    f"({len(col_c_rows)} en columna VALOR, {row_range})"
                )

            # 2. Crear nuevo workbook solo con Log_book (todo valores, sin fórmulas)
            wb_exp = openpyxl.Workbook()
            ws_exp = wb_exp.active
            ws_exp.title = SHEET_NAME

            for row in ws_src.iter_rows():
                for cell in row:
                    new_cell = ws_exp.cell(row=cell.row, column=cell.column)
                    # Convertir error Excel ("#VALUE!", "#REF!", etc.) a cadena vacía
                    val = cell.value
                    if isinstance(val, str) and val.startswith("#"):
                        val = ""   # error Excel → vacío provisional
                    # Fallback para celdas Stock:
                    # openpyxl devuelve None o "#VALUE!" para estas celdas;
                    # en ambos casos usamos el nombre extraído del richData XML
                    if not val and (cell.row, cell.column) in rich_values:
                        val = rich_values[(cell.row, cell.column)]
                    new_cell.value = val
                    # Copiar estilo completo: fuente, relleno, bordes, alineación, formato
                    if cell.has_style:
                        new_cell.font      = copy(cell.font)
                        new_cell.fill      = copy(cell.fill)
                        new_cell.border    = copy(cell.border)
                        new_cell.alignment = copy(cell.alignment)
                        new_cell.number_format = cell.number_format

            # Auto-ajustar anchos de columna según el contenido real
            # (evita "########" cuando el texto no cabe con el ancho del origen)
            import datetime as _dt
            for col_cells in ws_exp.columns:
                max_len = 0
                col_letter = get_column_letter(col_cells[0].column)
                for cell in col_cells:
                    if cell.value is None:
                        continue
                    if isinstance(cell.value, (_dt.datetime, _dt.date)):
                        cell_len = 12          # "DD/MM/YYYY" + margen
                    elif isinstance(cell.value, float):
                        cell_len = len(f"{cell.value:,.2f}") + 3
                    elif isinstance(cell.value, int):
                        cell_len = len(str(cell.value)) + 2
                    else:
                        cell_len = len(str(cell.value))
                    max_len = max(max_len, cell_len)
                ws_exp.column_dimensions[col_letter].width = min(max(max_len + 2, 8), 80)

            # Copiar altos de fila
            for row_num, row_dim in ws_src.row_dimensions.items():
                if row_dim.height:
                    ws_exp.row_dimensions[row_num].height = row_dim.height

            # Copiar celdas combinadas
            for merge_range in ws_src.merged_cells.ranges:
                try:
                    ws_exp.merge_cells(str(merge_range))
                except Exception:
                    pass

            # Serializar
            buf = io.BytesIO()
            wb_exp.save(buf)
            excel_bytes = buf.getvalue()

            # 3. Subir a Dropbox (carpeta DOCUMENTOS PUBLICADOS EN DISCORD)
            token = await self._get_access_token()
            if token:
                path_root = await self._get_path_root_header(token)
                headers_up = {
                    "Authorization":   f"Bearer {token}",
                    "Dropbox-API-Arg": json.dumps({
                        "path":       dest_dropbox_path,
                        "mode":       "overwrite",
                        "autorename": False,
                        "mute":       True,
                    }),
                    "Content-Type": "application/octet-stream",
                    **path_root,
                }
                async with aiohttp.ClientSession() as s:
                    async with s.post(
                        "https://content.dropboxapi.com/2/files/upload",
                        headers=headers_up,
                        data=excel_bytes,
                    ) as resp:
                        if resp.status != 200:
                            text = await resp.text()
                            logger.warning(f"Logbook export: Dropbox upload warning {resp.status}: {text[:150]}")
                        else:
                            logger.info(f"Logbook export: subido a Dropbox → {filename}")

            # 4. Publicar en Discord como adjunto
            # Usamos multipart manual para que el nombre del archivo no se URL-encode
            content_msg = (
                f"📊 **EXCEL OPERACIONES  —  {fecha.strftime('%d/%m/%Y')}**\n"
                f"Log_book completo con todos los valores actualizados."
            )
            boundary = "DiscordFileBoundary7MA4YW"
            nl = b"\r\n"
            discord_body = (
                b"--" + boundary.encode() + nl
                + b'Content-Disposition: form-data; name="payload_json"' + nl
                + b"Content-Type: application/json" + nl + nl
                + json.dumps({"content": content_msg}).encode()
                + nl
                + b"--" + boundary.encode() + nl
                + b'Content-Disposition: form-data; name="file"; filename="'
                + filename.encode("utf-8") + b'"' + nl
                + b"Content-Type: application/vnd.openxmlformats-officedocument"
                  b".spreadsheetml.sheet" + nl + nl
                + excel_bytes
                + nl
                + b"--" + boundary.encode() + b"--" + nl
            )
            async with aiohttp.ClientSession() as s:
                async with s.post(
                    discord_webhook,
                    data=discord_body,
                    headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
                ) as resp:
                    if resp.status in (200, 204):
                        logger.info(f"Logbook export: ✅ Excel de operaciones publicado ({filename})")
                        if self.log_channel:
                            await self.log_channel.send_info(
                                f"✅ Excel de operaciones publicado en Discord: **{filename}**"
                            )
                    else:
                        text = await resp.text()
                        raise RuntimeError(f"Discord webhook error {resp.status}: {text[:150]}")

        except Exception as e:
            logger.error(f"Logbook export error: {e}")
            if self.log_channel:
                await self.log_channel.send_error(
                    f"No se pudo exportar el Excel de operaciones: {e}"
                )
